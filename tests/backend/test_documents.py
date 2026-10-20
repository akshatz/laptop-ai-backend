import io
import uuid

import pytest
from docx import Document as DocxDocument


def minimal_pdf(text: str) -> bytes:
    """A one-page PDF with a text layer, built by hand (no PDF writer needed)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def docx_bytes(paragraph: str, cell: str) -> bytes:
    doc = DocxDocument()
    doc.add_paragraph(paragraph)
    doc.add_table(rows=1, cols=1).rows[0].cells[0].text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
async def chat_id(client, signed_in):
    return (await client.post("/chats", json={})).json()["id"]


async def test_upload_pdf_docx_and_text(client, chat_id, vector_store):
    files = [
        ("files", ("menu.pdf", minimal_pdf("Paneer tikka 250"), "application/pdf")),
        ("files", ("notes.docx", docx_bytes("Meeting on Friday", "Budget 5000"),
                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ("files", ("list.txt", b"milk, bread", "text/plain")),
    ]
    response = await client.post(f"/chats/{chat_id}/files", files=files)

    assert response.status_code == 201
    assert response.json() == {"documents_ingested": 3, "chunks_ingested": 3}
    texts = [d.page_content for d in vector_store.added]
    assert "Paneer tikka 250" in texts[0]
    assert "Meeting on Friday" in texts[1] and "Budget 5000" in texts[1]
    assert texts[2] == "milk, bread"
    assert {d.metadata["chat_id"] for d in vector_store.added} == {chat_id}
    assert [d.metadata["source"] for d in vector_store.added] == ["menu.pdf", "notes.docx", "list.txt"]


async def test_documents_lists_and_opens_your_own(client, chat_id, vector_store, sign_in, make_client):
    await client.post(f"/chats/{chat_id}/files", files=[("files", ("list.txt", b"milk, bread", "text/plain"))])
    await client.post(f"/chats/{chat_id}/documents", json={"documents": [{"text": "pasted words"}]})

    docs = (await client.get("/documents")).json()
    assert [d["filename"] for d in docs] == ["Pasted text 1.txt", "list.txt"]
    assert docs[1]["chat_id"] == chat_id and docs[1]["chat_title"] == "New Chat"
    assert docs[1]["size_bytes"] == 11 and docs[1]["chunks"] == 1 and "data" not in docs[1]

    opened = await client.get(f"/documents/{docs[1]['id']}/file")
    assert opened.status_code == 200 and opened.content == b"milk, bread"
    assert opened.headers["content-disposition"].startswith("inline;")
    assert opened.headers["x-content-type-options"] == "nosniff"

    async with await make_client() as other:
        await sign_in(other, "meera@example.com", "Meera")
        assert (await other.get("/documents")).json() == []
        assert (await other.get(f"/documents/{docs[1]['id']}/file")).status_code == 404
        response = await other.post(f"/chats/{chat_id}/files", files=[("files", ("x.txt", b"x", "text/plain"))])
        assert response.status_code == 404


async def test_html_is_served_as_plain_text(client, chat_id, vector_store):
    await client.post(f"/chats/{chat_id}/files", files=[("files", ("page.html", b"<script>alert(1)</script>", "text/html"))])
    doc_id = (await client.get("/documents")).json()[0]["id"]
    opened = await client.get(f"/documents/{doc_id}/file")
    assert opened.headers["content-type"].startswith("text/plain")


async def test_unsupported_or_empty_files_add_nothing(client, chat_id, vector_store):
    good = ("files", ("ok.txt", b"fine", "text/plain"))
    response = await client.post(f"/chats/{chat_id}/files", files=[good, ("files", ("photo.jpg", b"\xff\xd8", "image/jpeg"))])
    assert response.status_code == 400 and "unsupported file type" in response.json()["detail"]
    response = await client.post(f"/chats/{chat_id}/files", files=[("files", ("empty.txt", b"  ", "text/plain"))])
    assert response.status_code == 400 and "no text found" in response.json()["detail"]
    response = await client.post(f"/chats/{chat_id}/files", files=[("files", ("bad.pdf", b"not a pdf", "application/pdf"))])
    assert response.status_code == 400
    assert vector_store.added == [] and (await client.get("/documents")).json() == []


async def test_upload_limits(client, chat_id, vector_store, monkeypatch):
    from routers import chats

    six = [("files", (f"{i}.txt", b"x", "text/plain")) for i in range(6)]
    assert (await client.post(f"/chats/{chat_id}/files", files=six)).status_code == 400
    monkeypatch.setattr(chats, "MAX_FILE_BYTES", 10)
    big = [("files", ("big.txt", b"x" * 11, "text/plain"))]
    assert (await client.post(f"/chats/{chat_id}/files", files=big)).status_code == 413


async def test_milvus_failure_keeps_nothing(client, chat_id, vector_store):
    vector_store.add_error = RuntimeError("milvus down")
    response = await client.post(f"/chats/{chat_id}/files", files=[("files", ("a.txt", b"text", "text/plain"))])
    assert response.status_code == 502
    assert (await client.get("/documents")).json() == []


async def test_documents_need_a_session(client):
    assert (await client.get("/documents")).status_code == 401
    assert (await client.get(f"/documents/{uuid.uuid4()}/file")).status_code == 401
