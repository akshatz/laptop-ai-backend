"""Text from uploaded files, for adding them to a chat as context: PDF (pypdf, text layer only,
no OCR), Word .docx (python-docx, paragraphs and tables) and plain text (.txt, .md, .csv, .json,
.html, …). Raises UnsupportedFile for anything else, and for files with no text (e.g. a scanned
PDF without a text layer)."""

import io
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader
from pypdf.errors import PdfReadError

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".html", ".htm", ".xml", ".log", ".yaml", ".yml"}
SUPPORTED = sorted(TEXT_SUFFIXES | {".pdf", ".docx"})


class UnsupportedFile(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except (PdfReadError, ValueError) as exc:
            raise UnsupportedFile(f"{filename}: not a readable PDF ({exc})") from exc
    elif suffix == ".docx":
        try:
            doc = DocxDocument(io.BytesIO(data))
        except Exception as exc:  # python-docx raises several unrelated types for bad files
            raise UnsupportedFile(f"{filename}: not a readable Word file") from exc
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif suffix in TEXT_SUFFIXES:
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnsupportedFile(f"{filename}: unsupported file type (use {', '.join(SUPPORTED)})")

    text = text.strip()
    if not text:
        raise UnsupportedFile(f"{filename}: no text found (a scanned PDF needs OCR, which isn't supported here)")
    return text
